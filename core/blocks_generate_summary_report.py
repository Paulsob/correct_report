import os
import sys
import json
import pandas as pd
from datetime import datetime
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

# 📍 Определяем корень проекта относительно расположения скрипта
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
sys.path.append(PROJECT_ROOT)

import config

RU_DAYS = {0: "Пн", 1: "Вт", 2: "Ср", 3: "Чт", 4: "Пт", 5: "Сб", 6: "Вс"}
TARGET_SCHEDULES = ["4x2", "5x2h"]


def calculate_hours(time_range_str):
    """Надёжный парсер времени: работает с 'HH:MM' и 'HH:MM:SS'"""
    try:
        start_raw, end_raw = time_range_str.split(" - ")
        # Обрезаем секунды, если они есть, оставляем только HH:MM
        start = start_raw.strip()[:5]
        end = end_raw.strip()[:5]

        t1 = datetime.strptime(start, "%H:%M")
        t2 = datetime.strptime(end, "%H:%M")

        delta = (t2 - t1).total_seconds() / 3600.0
        if delta < 0:
            delta += 24.0
        return round(delta, 2)
    except Exception as e:
        print(f"⚠️ Ошибка парсинга времени '{time_range_str}': {e}")
        return 0.0


def generate_excel():
    prefix = "_".join(TARGET_SCHEDULES)
    print(f"📊 Загрузка данных для формирования отчета: [{prefix}]...")

    YEAR = int(config.TARGET_YEAR)
    MONTH = int(config.TARGET_MONTH)

    # 📂 Строим пути относительно корня проекта
    output_dir = os.path.join(PROJECT_ROOT, "output_data", str(YEAR), f"{MONTH:02d}", "4x2_5x2h")
    json_filename = f"{prefix}_BLOCK_final_schedule.json"
    excel_filename = f"{prefix}_BLOCK_drivers_report.xlsx"

    schedule_path = os.path.join(output_dir, json_filename)
    report_path = os.path.join(output_dir, excel_filename)

    # Безопасно резолвим путь к справочнику водителей
    drivers_path = config.DRIVERS_PREPARED
    if not os.path.isabs(drivers_path):
        drivers_path = os.path.join(PROJECT_ROOT, drivers_path)

    os.makedirs(output_dir, exist_ok=True)

    if not os.path.exists(schedule_path):
        print(f"❌ Ошибка: Файл расписания не найден: {schedule_path}")
        return
    if not os.path.exists(drivers_path):
        print(f"❌ Ошибка: Файл водителей не найден: {drivers_path}")
        return

    with open(schedule_path, "r", encoding="utf-8") as f:
        schedule_data = json.load(f)
    with open(drivers_path, "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    # 1. Инициализация структуры водителей
    drivers_stats = {}
    for driver in drivers_data.get("drivers", []):
        if driver.get("schedule_type") not in TARGET_SCHEDULES:
            continue

        tab = str(driver.get("tab_number"))
        drivers_stats[tab] = {
            "График": driver.get("schedule_type"),
            "Табельный номер": tab,
            "Смен в месяц": 0,
            "Часов в месяц": 0.0
        }

    # 🔍 Находим максимальный день месяца из расписания
    max_day_in_month = 0
    for day_info in schedule_data:
        date_str = day_info.get("дата", "")
        if date_str:
            day = datetime.strptime(date_str, "%Y-%m-%d").day
            if day > max_day_in_month:
                max_day_in_month = day

    if max_day_in_month == 0:
        print("⚠️ В расписании нет дат. Отчёт не может быть сформирован.")
        return

    # 🛡️ Гарантированно создаём ячейки для всех дней (1..max_day)
    for tab, data in drivers_stats.items():
        original_days = next(
            (d.get("days", {}) for d in drivers_data.get("drivers", []) if str(d.get("tab_number")) == tab),
            {}
        )
        for d in range(1, max_day_in_month + 1):
            day_str = str(d)
            val = original_days.get(day_str, "Р")  # По умолчанию "Р" (РЕЗЕРВ), если в справочнике нет
            data[day_str] = "Выходной" if val == "В" else "РЕЗЕРВ"

    # 2. Накладываем фактические назначения
    print("⚙️ Обработка назначений на смены...")
    for day_info in schedule_data:
        date_str = day_info.get("дата", "")
        if not date_str:
            continue

        day_number = datetime.strptime(date_str, "%Y-%m-%d").day
        day_str = str(day_number)

        for assignment in day_info.get("успешные_назначения", []):
            tab_number = str(assignment.get("таб_номер_водителя"))
            if tab_number not in drivers_stats:
                continue

            time_str = assignment.get("время", "")
            hours = calculate_hours(time_str)

            drivers_stats[tab_number]["Смен в месяц"] += 1
            drivers_stats[tab_number]["Часов в месяц"] += hours

            cell_text = (f"Маршрут: {assignment.get('маршрут', '-')}\n"
                         f"Режим: {assignment.get('смена', '-')}\n"
                         f"Смена №: {assignment.get('трамвай', '-')}\n"
                         f"{hours} часов")

            current_val = drivers_stats[tab_number][day_str]
            if current_val in ["РЕЗЕРВ", "Выходной"]:
                drivers_stats[tab_number][day_str] = cell_text
            else:
                drivers_stats[tab_number][day_str] += f"\n\n{cell_text}"

    # 3. ФИЛЬТРАЦИЯ: Оставляем только тех, кто работал
    drivers_stats = {tab: data for tab, data in drivers_stats.items() if data["Смен в месяц"] > 0}
    data_for_df = list(drivers_stats.values())

    if not data_for_df:
        print("ℹ️ Нет ни одного водителя с назначенными сменами. Отчет пуст.")
        return

    df = pd.DataFrame(data_for_df)
    df = df.sort_values(by=["График", "Табельный номер"])
    df["Часов в месяц"] = df["Часов в месяц"].round(2)

    # 4. Умная шапка
    day_cols = [str(d) for d in range(1, max_day_in_month + 1)]
    base_cols = ["График", "Табельный номер", "Смен в месяц", "Часов в месяц"]
    columns_to_keep = base_cols + day_cols

    # Безопасный выбор колонок
    df = df[[c for c in columns_to_keep if c in df.columns]]

    new_columns = []
    for col in df.columns:
        if col in base_cols:
            new_columns.append(col)
        else:
            day = int(col)
            current_date = datetime(YEAR, MONTH, day)
            new_columns.append(f"{day}\n{RU_DAYS[current_date.weekday()]}")

    df.columns = new_columns

    # 5. Экспорт и форматирование
    print("💾 Сохранение и форматирование в Excel...")
    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    gray_fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    center_aligned_text = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="center")

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        sheet_name = f"Отчет {prefix}"[:31]
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        # 📐 Настройки листа
        ws.freeze_panes = "E2"
        ws.row_dimensions[1].height = 30
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 65

        # 📏 Ширина колонок
        for idx, col in enumerate(df.columns, start=1):
            col_letter = get_column_letter(idx)
            ws.column_dimensions[col_letter].width = 12 if idx <= 4 else 16

        # 🎨 Оформление ячеек дней (мин. колонка 5)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=5, max_col=ws.max_column):
            for cell in row:
                cell.alignment = center_aligned_text
                if cell.value == "РЕЗЕРВ":
                    cell.fill = yellow_fill
                elif cell.value == "Выходной":
                    cell.fill = gray_fill
                    cell.font = Font(color="000000")

        # 📊 Оформление первых 4 колонок
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.alignment = center_alignment

        # 🔠 Шапка
        for cell in ws[1]:
            cell.alignment = center_aligned_text
            cell.font = Font(bold=True)

    print(f"✅ Отчет успешно сформирован:\n{report_path}")


if __name__ == "__main__":
    generate_excel()
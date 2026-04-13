import os
import json
from openpyxl import load_workbook

# -------------------------------
# Пути
# -------------------------------
BASE_DIR = r"C:\Users\psobo\PycharmProjects\correct_report"

INPUT_FILE = os.path.join(BASE_DIR, "input_data", "permits.xlsx")
OUTPUT_FILE = os.path.join(BASE_DIR, "input_data", "prepared_data", "permits_real.json")

# -------------------------------
# Загрузка Excel
# -------------------------------
wb = load_workbook(INPUT_FILE, data_only=True)
ws = wb["Лист_1"]

# -------------------------------
# Обработка данных
# -------------------------------
result = {}

for row in ws.iter_rows(min_row=2):  # пропускаем заголовок
    tab_number = row[1].value  # столбец B
    model = row[3].value       # столбец D
    status = row[4].value      # столбец E

    # Фильтр: только "Нет"
    if status != "Нет":
        continue

    if tab_number is None or model is None:
        continue

    tab_number = str(tab_number).strip().replace(" ", "")

    if tab_number not in result:
        result[tab_number] = []

    # Убираем только точные дубликаты, сохраняя порядок
    if model not in result[tab_number]:
        result[tab_number].append(model)

# -------------------------------
# Преобразование в список
# -------------------------------
output_data = [
    {
        "tab_number": tab_number,
        "models": models
    }
    for tab_number, models in result.items()
]

# -------------------------------
# Сохранение JSON
# -------------------------------
os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)

with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
    json.dump(output_data, f, ensure_ascii=False, indent=2)

print("Готово! JSON сохранён в:", OUTPUT_FILE)
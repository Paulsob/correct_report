from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import range_boundaries


# -------------------------------
# Константы структуры Excel
# -------------------------------
BLOCK_WIDTH = 12
PS_OFFSET = 1
SHIFT1_START_OFFSET = 5
SHIFT1_END_OFFSET = 6
SHIFT2_START_OFFSET = 10
SHIFT2_END_OFFSET = 11


# -------------------------------
# Вспомогательные функции
# -------------------------------
def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    text = str(value).strip()
    return text if text else None


def is_record_start(value: Any) -> bool:
    text = normalize_text(value)
    if not text:
        return False
    if text.startswith("Маршрут"):
        return False
    if text in {"ЗАПАС"}:
        return False
    return bool(re.search(r"\d", text))


def parse_route_number(route_title: str):
    m = re.search(r"Маршрут\s+(\d+)", route_title)
    return int(m.group(1)) if m else None


def parse_schedule_title(ws):
    title = ""
    for row in range(1, 6):
        for col in range(1, 20):
            value = ws.cell(row, col).value
            if isinstance(value, str) and re.search(r"\d{2}\.\d{2}\.\d{4}", value):
                title = value.strip()
                break
        if title:
            break

    weekday_match = re.search(r"\(([^)]+)\)", title)
    weekday = weekday_match.group(1).strip().lower() if weekday_match else ""

    day_type = "выходной" if weekday in {"суббота", "воскресенье"} else "рабочий"
    return day_type


# -------------------------------
# Поиск маршрутов
# -------------------------------
def collect_route_headers(ws):
    headers = []

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        value = ws.cell(min_row, min_col).value

        if isinstance(value, str) and value.startswith("Маршрут"):
            headers.append({
                "row": min_row,
                "col": min_col,
                "title": value.strip(),
            })

    headers.sort(key=lambda x: (x["row"], x["col"]))
    return headers


# -------------------------------
# Парсинг блока маршрута
# -------------------------------
def parse_route_block(ws, header, next_header_row, day_type):
    start_col = header["col"]
    start_row = header["row"] + 1
    stop_row = next_header_row - 1 if next_header_row else ws.max_row

    route_number = parse_route_number(header["title"])
    trams = []

    row = start_row
    while row <= stop_row:
        trip_value = ws.cell(row, start_col).value

        if is_record_start(trip_value):
            tram = {
                "номер": normalize_text(trip_value),
                "номер_пс": normalize_text(ws.cell(row, start_col + PS_OFFSET).value),
                "смена_1": {
                    "отправление": normalize_text(ws.cell(row, start_col + SHIFT1_START_OFFSET).value),
                    "прибытие": normalize_text(ws.cell(row, start_col + SHIFT1_END_OFFSET).value),
                },
                "смена_2": {
                    "отправление": normalize_text(ws.cell(row, start_col + SHIFT2_START_OFFSET).value),
                    "прибытие": normalize_text(ws.cell(row, start_col + SHIFT2_END_OFFSET).value),
                },
            }
            trams.append(tram)

        row += 1

    return {
        "маршрут": route_number,
        "день": day_type,
        "трамваи": trams,
    }


# -------------------------------
# Основной парсер
# -------------------------------
def parse_workbook(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    day_type = parse_schedule_title(ws)
    headers = collect_route_headers(ws)

    blocks = []

    for i, header in enumerate(headers):
        next_header_row = None

        for j in range(i + 1, len(headers)):
            if headers[j]["col"] == header["col"]:
                next_header_row = headers[j]["row"]
                break

        block = parse_route_block(ws, header, next_header_row, day_type)
        blocks.append(block)

    return blocks


# -------------------------------
# MAIN
# -------------------------------
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", default="040426.xlsx")
    args = parser.parse_args()

    # 📍 ВАЖНО: правильные пути относительно файла скрипта
    script_dir = Path(__file__).resolve().parent
    blocks_dir = script_dir.parent

    input_file = blocks_dir / "input_data" / args.file
    output_file = blocks_dir / "input_data" / "raw_data" / "new_raw_schedule_weekend.json"

    # Проверка
    if not input_file.exists():
        raise FileNotFoundError(f"Файл не найден: {input_file}")

    data = parse_workbook(input_file)

    output_file.parent.mkdir(parents=True, exist_ok=True)

    with output_file.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"Готово: {output_file}")


if __name__ == "__main__":
    main()
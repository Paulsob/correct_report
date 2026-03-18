import json
import os
from datetime import datetime, timedelta
from ortools.sat.python import cp_model
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

# Настройки
MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4,
    "май": 5, "июнь": 6, "июль": 7, "август": 8,
    "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12
}


class Shift:
    def __init__(self, shift_id, day, type_str, tram_num, start_dt, end_dt):
        self.id = shift_id
        self.day = day
        self.type_str = type_str
        self.tram_num = tram_num
        self.start_dt = start_dt
        self.end_dt = end_dt
        self.duration = (end_dt - start_dt).total_seconds() / 3600.0


def parse_time(base_date, time_str):
    h, m = map(int, time_str.split(':'))
    dt = base_date.replace(hour=h, minute=m, second=0, microsecond=0)
    return dt


def generate_excel_report(assigned_shifts, all_shifts, covered_shift_ids, drivers_dict, max_days,
                          filename="Отчет_Расписание.xlsx"):
    print(f"\nГенерация Excel-отчета: {filename}...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расписание"

    # --- СТИЛИ ---
    bold_font = Font(bold=True)
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                         bottom=Side(style='thin'))

    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Успешная смена
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Резерв
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")  # Выходной
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Незакрытая смена

    # --- ЗАГОЛОВКИ ---
    headers = ["Таб. №", "Смен", "Часов"] + [str(d) for d in range(1, max_days + 1)]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = bold_font
        cell.alignment = center_aligned
        cell.fill = header_fill
        cell.border = thin_border

    # --- ЗАПОЛНЕНИЕ ВОДИТЕЛЕЙ ---
    # Сортируем водителей по количеству часов (по убыванию)
    drivers_sorted = sorted(assigned_shifts.items(), key=lambda item: sum(s.duration for s in item[1]), reverse=True)

    for d_id, shifts_list in drivers_sorted:
        if not shifts_list:
            continue  # Пропускаем водителей, которых алгоритм вообще не вызвал в этом месяце

        shifts_list.sort(key=lambda s: s.start_dt)
        total_shifts = len(shifts_list)
        total_hours = sum(s.duration for s in shifts_list)

        row_cells = [d_id, total_shifts, round(total_hours, 1)]
        ws.append(row_cells)
        current_row = ws.max_row

        # Применяем базовый стиль к первым 3 колонкам
        for col_idx in range(1, 4):
            ws.cell(row=current_row, column=col_idx).alignment = center_aligned
            ws.cell(row=current_row, column=col_idx).border = thin_border

        shifts_by_day = {s.day: s for s in shifts_list}
        driver_pattern = drivers_dict[d_id]["pattern"]

        # Заполняем дни
        for day in range(1, max_days + 1):
            col_idx = day + 3
            cell = ws.cell(row=current_row, column=col_idx)
            cell.alignment = center_aligned
            cell.border = thin_border

            pattern_val = driver_pattern.get(day, "В")

            if day in shifts_by_day:
                # ВОДИТЕЛЬ РАБОТАЕТ (ЗЕЛЕНЫЙ)
                current_shift = shifts_by_day[day]
                idx = shifts_list.index(current_shift)
                rest_str = "---"
                if idx > 0:
                    prev_shift = shifts_list[idx - 1]
                    rest_hours = (current_shift.start_dt - prev_shift.end_dt).total_seconds() / 3600.0
                    rest_str = f"{round(rest_hours, 1)}ч"

                cell.value = f"Р: {round(current_shift.duration, 1)}ч\nО: {rest_str}"
                cell.fill = green_fill
            else:
                if pattern_val in ["1", "2"]:
                    # ДОЛЖЕН БЫЛ РАБОТАТЬ, НО НЕ ПОЛУЧИЛ СМЕНУ (ЖЕЛТЫЙ - РЕЗЕРВ)
                    cell.value = "РЕЗЕРВ"
                    cell.fill = yellow_fill
                else:
                    # ЗАКОННЫЙ ВЫХОДНОЙ (СЕРЫЙ)
                    cell.value = "В"
                    cell.fill = gray_fill
                    cell.font = Font(color="808080")

    # --- ИНФОРМАЦИЯ О НЕЗАКРЫТЫХ СМЕНАХ (ПОДВАЛ) ---
    uncovered_by_day = {day: [] for day in range(1, max_days + 1)}
    for s in all_shifts:
        if s.id not in covered_shift_ids:
            uncovered_by_day[s.day].append(f"См.{s.type_str} (Тр.{s.tram_num})")

    ws.append([])  # Пустая строка для отступа
    bottom_row_idx = ws.max_row + 1
    ws.cell(row=bottom_row_idx, column=1).value = "НЕЗАКРЫТЫЕ СМЕНЫ:"
    ws.cell(row=bottom_row_idx, column=1).font = bold_font

    for day in range(1, max_days + 1):
        cell = ws.cell(row=bottom_row_idx, column=day + 3)
        cell.alignment = center_aligned
        cell.border = thin_border
        uncovered = uncovered_by_day[day]

        if uncovered:
            cell.value = "\n".join(uncovered)
            cell.fill = red_fill
            cell.font = Font(color="9C0006", bold=True)
        else:
            cell.value = "ОК"
            cell.font = Font(color="006100")
            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    # --- ШИРИНА КОЛОНОК ---
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 8
    for col in range(4, max_days + 4):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14

    wb.save(filename)
    print("Готово! Отчет сохранен.")


def solve_schedule():
    if not os.path.exists("10_drivers_october.json") or not os.path.exists("schedule.json"):
        print("ОШИБКА: Файлы 10_drivers_october.json и schedule.json не найдены!")
        return

    with open("10_drivers_october.json", "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    with open("schedule.json", "r", encoding="utf-8") as f:
        schedule_data = json.load(f)

    year = drivers_data.get("year", 2026)
    month_name = drivers_data.get("month", "Октябрь").lower()
    month = MONTHS.get(month_name, 10)

    schedule_by_type = {item["день"]: item["трамваи"] for item in schedule_data}

    max_days = 0
    drivers_dict = {}
    drivers = []

    for d in drivers_data["drivers"]:
        pattern = {item["day"]: item["value"] for item in d.get("days", [])}
        if pattern:
            max_days = max(max_days, max(pattern.keys()))
        driver_obj = {"tab_number": d["tab_number"], "pattern": pattern}
        drivers.append(driver_obj)
        drivers_dict[d["tab_number"]] = driver_obj

    shifts = []
    shift_id_counter = 0

    # ГЕНЕРАЦИЯ СМЕН
    for day in range(1, max_days + 1):
        current_date = datetime(year, month, day)
        day_type = "рабочий" if current_date.weekday() < 5 else "выходной"

        for tram in schedule_by_type.get(day_type, []):
            tram_num = tram["номер"]

            s1_start = parse_time(current_date, tram["смена_1"]["отправление"])
            s1_end = parse_time(current_date, tram["смена_1"]["прибытие"])
            if s1_end < s1_start: s1_end += timedelta(days=1)
            shifts.append(Shift(shift_id_counter, day, "1", tram_num, s1_start, s1_end))
            shift_id_counter += 1

            s2_start = parse_time(current_date, tram["смена_2"]["отправление"])
            s2_end = parse_time(current_date, tram["смена_2"]["прибытие"])
            if s2_end < s2_start: s2_end += timedelta(days=1)
            shifts.append(Shift(shift_id_counter, day, "2", tram_num, s2_start, s2_end))
            shift_id_counter += 1

    model = cp_model.CpModel()
    x = {}

    # Переменные только для совпадений графика и смен (ЖЕСТКОЕ ПРАВИЛО: ВЫХОДНЫЕ ЗАПРЕЩЕНЫ)
    for shift in shifts:
        for driver in drivers:
            d_id = driver["tab_number"]
            if driver["pattern"].get(shift.day) == shift.type_str:
                x[(d_id, shift.id)] = model.NewBoolVar(f"x_{d_id}_s_{shift.id}")

    # ОГРАНИЧЕНИЕ 1: Максимум 1 водитель на смену (позволяем оставлять пустые)
    is_covered = {}
    for shift in shifts:
        shift_vars = [x[(d["tab_number"], shift.id)] for d in drivers if (d["tab_number"], shift.id) in x]
        is_covered[shift.id] = model.NewBoolVar(f"covered_{shift.id}")
        # Если сумма переменных = 1, значит смена закрыта (is_covered = 1)
        model.Add(sum(shift_vars) == is_covered[shift.id])

    # ОГРАНИЧЕНИЕ 2: Отдых
    for driver in drivers:
        d_id = driver["tab_number"]
        driver_shifts = [s for s in shifts if (d_id, s.id) in x]

        for i in range(len(driver_shifts)):
            for j in range(i + 1, len(driver_shifts)):
                s1 = driver_shifts[i]
                s2 = driver_shifts[j]

                first, second = (s1, s2) if s1.start_dt < s2.start_dt else (s2, s1)

                gap_hours = (second.start_dt - first.end_dt).total_seconds() / 3600.0
                required_rest = first.duration * 2.0

                if gap_hours < required_rest:
                    model.Add(x[(d_id, first.id)] + x[(d_id, second.id)] <= 1)

    # ЦЕЛИ ОПТИМИЗАЦИИ
    is_used = {}
    for driver in drivers:
        d_id = driver["tab_number"]
        is_used[d_id] = model.NewBoolVar(f"used_{d_id}")
        driver_vars = [x[(d_id, s.id)] for s in shifts if (d_id, s.id) in x]
        if driver_vars:
            model.AddMaxEquality(is_used[d_id], driver_vars)
        else:
            model.Add(is_used[d_id] == 0)

    # Формула: Максимизируем закрытые смены (приоритет №1), Минимизируем кол-во нанятых людей (приоритет №2)
    # Это автоматически заставляет алгоритм не плодить "РЕЗЕРВ" для работающих водителей
    model.Maximize(1000 * sum(is_covered.values()) - sum(is_used.values()))

    print("Поиск оптимального расписания (OR-Tools)...")
    solver = cp_model.CpSolver()
    status = solver.Solve(model)

    if status == cp_model.OPTIMAL or status == cp_model.FEASIBLE:
        covered_count = sum(1 for s in shifts if solver.BooleanValue(is_covered[s.id]))
        used_count = int(sum(solver.BooleanValue(is_used[d["tab_number"]]) for d in drivers))

        print("\n=== РАСПИСАНИЕ СОСТАВЛЕНО ===")
        print(f"Закрыто смен: {covered_count} из {len(shifts)}")
        print(f"Задействовано водителей: {used_count} из {len(drivers)}")

        assigned_shifts = {d["tab_number"]: [] for d in drivers}
        covered_shift_ids = set()

        for shift in shifts:
            if solver.BooleanValue(is_covered[shift.id]):
                covered_shift_ids.add(shift.id)
            for driver in drivers:
                d_id = driver["tab_number"]
                if (d_id, shift.id) in x and solver.BooleanValue(x[(d_id, shift.id)]):
                    assigned_shifts[d_id].append(shift)

        report_name = f"Расписание_{month_name.capitalize()}_{year}.xlsx"
        generate_excel_report(assigned_shifts, shifts, covered_shift_ids, drivers_dict, max_days, report_name)

    else:
        print("\n[КРИТИЧЕСКАЯ ОШИБКА] Невозможно составить расписание даже частично.")


if __name__ == "__main__":
    solve_schedule()
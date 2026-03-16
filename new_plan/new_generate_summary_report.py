import json
from datetime import datetime
from pathlib import Path
from collections import Counter, defaultdict
from typing import List, Dict, Any, Union
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment


def summarize(schedule_data: Union[List, Dict]) -> Dict[str, Any]:
    total_days = 0
    total_shifts = 0
    uncovered_shifts = 0

    driver_shifts = Counter()
    driver_hours = defaultdict(float)

    actual_work_matrix = defaultdict(dict)

    # НОВОЕ: Таймлайн для расчета отдыха
    driver_timeline = defaultdict(list)

    if isinstance(schedule_data, dict):
        days_iterable = schedule_data.items()
    else:
        days_iterable = enumerate(schedule_data, start=1)

    for day_num, day_data in days_iterable:
        day_str = str(day_num)
        total_days = max(total_days, int(day_num))

        if isinstance(day_data, dict):
            if "трамваи" in day_data:
                trams = day_data["трамваи"]
            elif "roster" in day_data:
                trams = day_data["roster"]
            else:
                trams = [{"номер": k, **v} for k, v in day_data.items()]
        else:
            trams = day_data

        for tram in trams:
            for key in ("смена_1", "смена_2", "shift_1", "shift_2", "1", "2"):
                shift = tram.get(key)
                if shift is not None:
                    total_shifts += 1
                    driver_id = None
                    hours = 0.0

                    if isinstance(shift, dict):
                        driver_id = shift.get("driver_tab", shift.get("driver", shift.get("driver_id")))
                        hours = float(shift.get("work_hours", shift.get("duration", 0.0)))

                        # Собираем даты начала и конца для расчета отдыха
                        s_start = shift.get("start_dt")
                        s_end = shift.get("end_dt")
                        if driver_id and s_start and s_end:
                            driver_timeline[str(driver_id)].append({
                                "day": day_str,
                                "start": datetime.fromisoformat(s_start),
                                "end": datetime.fromisoformat(s_end)
                            })

                    elif isinstance(shift, str):
                        if "Водитель" in shift:
                            parts = shift.split()
                            if len(parts) >= 2:
                                driver_id = parts[1]
                                try:
                                    hours = float(parts[2].replace('(', '').replace('ч)', '').replace('ч', ''))
                                except:
                                    hours = 8.0
                        else:
                            driver_id = shift

                    if driver_id and str(driver_id).strip() not in ["", "None", "НЕТ ВОДИТЕЛЯ", "NO_DRIVER"]:
                        d_id_str = str(driver_id).strip()
                        driver_shifts[d_id_str] += 1
                        driver_hours[d_id_str] += hours
                        actual_work_matrix[d_id_str][day_str] = hours
                    else:
                        uncovered_shifts += 1

    unique_drivers = len(driver_shifts)

    # НОВОЕ: Расчет отдыха для каждой смены
    driver_daily_rest = defaultdict(dict)
    for d_id, shifts in driver_timeline.items():
        # Сортируем смены водителя в хронологическом порядке
        shifts.sort(key=lambda x: x["start"])
        for i in range(1, len(shifts)):
            prev_shift = shifts[i - 1]
            curr_shift = shifts[i]
            # Считаем разницу между КОНЦОМ прошлой смены и НАЧАЛОМ текущей
            gap = (curr_shift["start"] - prev_shift["end"]).total_seconds() / 3600.0
            driver_daily_rest[d_id][curr_shift["day"]] = round(gap, 1)

    distribution = {
        "Переработка (>160ч)": [],
        "Норма (140-160ч)": [],
        "Недоработка (80-140ч)": [],
        "Огрызки (<80ч)": [],
        "Полный резерв (0ч)": []
    }

    sorted_active = sorted(driver_hours.items(), key=lambda x: x[1], reverse=True)

    for d_id, hours in sorted_active:
        if hours > 160:
            distribution["Переработка (>160ч)"].append(d_id)
        elif hours >= 140:
            distribution["Норма (140-160ч)"].append(d_id)
        elif hours >= 80:
            distribution["Недоработка (80-140ч)"].append(d_id)
        else:
            distribution["Огрызки (<80ч)"].append(d_id)

    covered_shifts = total_shifts - uncovered_shifts
    coverage_percent = (covered_shifts / total_shifts * 100) if total_shifts > 0 else 0

    return {
        "total_days": total_days,
        "total_shifts": total_shifts,
        "covered_shifts": covered_shifts,
        "uncovered_shifts": uncovered_shifts,
        "coverage_percent": coverage_percent,
        "unique_drivers_used": unique_drivers,
        "distribution": distribution,
        "driver_hours": driver_hours,
        "driver_shifts": driver_shifts,
        "actual_work_matrix": actual_work_matrix,
        "driver_daily_rest": driver_daily_rest  # Передаем отдых в Excel
    }


def generate_excel_report(summary: Dict[str, Any], original_drivers: List[Dict], filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Табель расписания"

    header_font = Font(bold=True)
    # Включаем wrap_text (перенос по словам), чтобы часы и отдых были на разных строчках в ячейке
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    fill_work = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_off = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    fill_reserve = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    fill_hole = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    days_in_month = summary["total_days"]

    headers = ["Таб. №", "Категория", "Часы", "Смены"] + [str(d) for d in range(1, days_in_month + 1)]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = center_align

    driver_patterns = {}
    for d in original_drivers:
        d_id = str(d["tab_number"])
        driver_patterns[d_id] = d.get("pattern", {})

    all_drivers_data = []

    for cat_name, d_ids in summary["distribution"].items():
        for d_id in d_ids:
            all_drivers_data.append({
                "id": d_id, "category": cat_name,
                "hours": summary["driver_hours"][d_id],
                "shifts": summary["driver_shifts"][d_id]
            })

    worked_ids = set(summary["driver_hours"].keys())
    for d_id in driver_patterns.keys():
        if d_id not in worked_ids:
            all_drivers_data.append({
                "id": d_id, "category": "Полный резерв (0ч)", "hours": 0.0, "shifts": 0
            })

    for d_data in all_drivers_data:
        d_id = d_data["id"]
        row = [d_id, d_data["category"], round(d_data["hours"], 1), d_data["shifts"]]

        work_days = summary["actual_work_matrix"].get(d_id, {})
        rest_days = summary["driver_daily_rest"].get(d_id, {})
        pattern = driver_patterns.get(d_id, {})

        cells_to_color = []

        for day in range(1, days_in_month + 1):
            day_str = str(day)
            worked_hours = work_days.get(day_str)
            scheduled_status = str(pattern.get(day, pattern.get(day_str, "В")))

            if worked_hours:
                rest = rest_days.get(day_str)
                # Если это не первый рабочий день, добавляем инфу об отдыхе с новой строки
                if rest is not None:
                    row.append(f"Р ({worked_hours}ч)\n[Отд: {rest}ч]")
                else:
                    row.append(f"Р ({worked_hours}ч)")
                cells_to_color.append(fill_work)
            else:
                if scheduled_status == "В":
                    row.append("В (Вых)")
                    cells_to_color.append(fill_off)
                else:
                    row.append("Резерв")
                    cells_to_color.append(fill_reserve)

        ws.append(row)

        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 30  # Увеличиваем высоту строки для двух строчек текста

        for col_idx, fill_color in enumerate(cells_to_color, start=5):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill_color
            cell.alignment = center_align

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 8
    for i in range(5, days_in_month + 5):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12

    wb.save(filepath)


def auto_generate_report(schedule_data, drivers_list, output_txt=None, output_excel=None):
    summary = summarize(schedule_data)

    print(f"\n=== КРАТКАЯ СВОДКА ===")
    print(f"Покрытие графика: {summary['coverage_percent']:.1f}%")
    print(f"Незакрытых смен (Дыры): {summary['uncovered_shifts']}")
    for cat, d_ids in summary["distribution"].items():
        print(f"[{cat}] - {len(d_ids)} чел.")

    if output_excel:
        try:
            generate_excel_report(summary, drivers_list, output_excel)
            print(f"\n[+] Подробный Excel-отчет сохранен в: {output_excel}")
        except Exception as e:
            print(f"[-] Не удалось сохранить Excel: {e}")
import sys
import os

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from blocks import config

import json
import pandas as pd
from datetime import datetime
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

RU_DAYS = {0: "Пн", 1: "Вт", 2: "Ср", 3: "Чт", 4: "Пт", 5: "Сб", 6: "Вс"}


def calculate_hours(time_range_str):
    try:
        start_str, end_str = time_range_str.split(" - ")
        time_format = "%H:%M"
        start_time = datetime.strptime(start_str, time_format)
        end_time = datetime.strptime(end_str, time_format)

        delta_hours = (end_time - start_time).total_seconds() / 3600.0
        if delta_hours < 0:
            delta_hours += 24.0

        return round(delta_hours, 2)
    except Exception as e:
        print(f"Ошибка парсинга времени '{time_range_str}': {e}")
        return 0.0


def generate_excel():
    print(f"Загрузка данных для графика: [{config.TARGET_SCHEDULE}]...")

    if not os.path.exists(config.FINAL_SCHEDULE) or not os.path.exists(config.DRIVERS_PREPARED):
        print("Файлы данных не найдены! Сначала запустите планировщик.")
        return

    with open(config.FINAL_SCHEDULE, "r", encoding="utf-8") as f:
        schedule_data = json.load(f)

    with open(config.DRIVERS_PREPARED, "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    # 1. Формируем "Подложку" (Берем только водителей нужного графика!)
    drivers_stats = {}
    for driver in drivers_data.get("drivers", []):
        if driver.get("schedule_type") != config.TARGET_SCHEDULE:
            continue

        tab = driver.get("tab_number")
        drivers_stats[tab] = {
            "Табельный номер": tab,
            "Смен в месяц": 0,
            "Часов в месяц": 0.0
        }

        # Читаем сжатый словарь дней ("1": "1", "2": "В")
        for day_str, val in driver.get("days", {}).items():
            if val == "В":
                drivers_stats[tab][day_str] = "Выходной"
            else:
                drivers_stats[tab][day_str] = "РЕЗЕРВ"

    max_day_in_month = 0
    report_year = int(config.TARGET_YEAR)
    report_month = int(config.TARGET_MONTH)

    # 2. Накладываем фактические назначения
    print("Обработка назначений на смены...")
    for day_info in schedule_data:
        date_str = day_info.get("дата", "")
        if not date_str:
            continue

        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        day_number = date_obj.day

        if day_number > max_day_in_month:
            max_day_in_month = day_number

        for assignment in day_info.get("успешные_назначения", []):
            tab_number = assignment.get("таб_номер_водителя")
            shift_type = assignment.get("смена")
            time_str = assignment.get("время")
            route_num = assignment.get("маршрут")
            tram_num = assignment.get("трамвай")

            if tab_number not in drivers_stats:
                continue

            hours = calculate_hours(time_str)

            drivers_stats[tab_number]["Смен в месяц"] += 1
            drivers_stats[tab_number]["Часов в месяц"] += hours

            cell_text = (f"Маршрут: {route_num}\n"
                         f"Режим: {shift_type}\n"
                         f"Смена №: {tram_num}\n"
                         f"{hours} часов")

            current_val = drivers_stats[tab_number][str(day_number)]
            if current_val in ["РЕЗЕРВ", "Выходной"]:
                drivers_stats[tab_number][str(day_number)] = cell_text
            else:
                drivers_stats[tab_number][str(day_number)] += f"\n\n{cell_text}"

    # 3. ФИЛЬТРАЦИЯ ВОДИТЕЛЕЙ (Оставляем только тех, кто работал)
    drivers_stats = {tab: data for tab, data in drivers_stats.items() if data["Смен в месяц"] > 0}

    data_for_df = list(drivers_stats.values())
    if not data_for_df:
        print(f"Нет ни одного водителя ({config.TARGET_SCHEDULE}) с назначенными сменами. Отчет пуст.")
        return

    df = pd.DataFrame(data_for_df)
    df = df.sort_values(by="Табельный номер")
    df["Часов в месяц"] = df["Часов в месяц"].round(2)

    # 4. Создаем умную шапку
    new_columns = []
    columns_to_keep = ["Табельный номер", "Смен в месяц", "Часов в месяц"] + [str(d) for d in
                                                                              range(1, max_day_in_month + 1)]
    df = df[columns_to_keep]

    for col in df.columns:
        if col in ["Табельный номер", "Смен в месяц", "Часов в месяц"]:
            new_columns.append(col)
        else:
            day = int(col)
            current_date = datetime(report_year, report_month, day)
            day_of_week_str = RU_DAYS[current_date.weekday()]
            new_columns.append(f"{day}\n{day_of_week_str}")

    df.columns = new_columns

    # 5. Сохранение и визуальное оформление
    print("Сохранение и форматирование в Excel...")

    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    gray_fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    center_aligned_text = Alignment(horizontal="center", vertical="center", wrap_text=True)

    with pd.ExcelWriter(config.REPORT_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=f"График {config.TARGET_SCHEDULE}")
        worksheet = writer.sheets[f"График {config.TARGET_SCHEDULE}"]

        worksheet.freeze_panes = "D2"
        worksheet.row_dimensions[1].height = 30

        for r in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[r].height = 65

        for idx, col in enumerate(df.columns):
            col_letter = get_column_letter(idx + 1)
            worksheet.column_dimensions[col_letter].width = 15 if idx < 3 else 16

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=4, max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = center_aligned_text
                if cell.value == "РЕЗЕРВ":
                    cell.fill = yellow_fill
                elif cell.value == "Выходной":
                    cell.fill = gray_fill
                    cell.font = Font(color="000000")

        for cell in worksheet[1]:
            cell.alignment = center_aligned_text
            cell.font = Font(bold=True)

    print(f"Отчет сформирован и сохранен в: {config.REPORT_EXCEL}")


if __name__ == "__main__":
    generate_excel()

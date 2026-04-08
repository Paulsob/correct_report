import os
import sys
import json
import calendar
from datetime import datetime
from collections import defaultdict

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
except ImportError:
    print("Ошибка: Не установлена библиотека openpyxl. Выполните 'pip install openpyxl'")
    sys.exit(1)

current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)
sys.path.append(parent_dir)

from another_idea import config


def format_duration(seconds):
    if seconds is None: return "00:00"
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    return f"{hours:02d}:{minutes:02d}"


def is_driver_absent(tab, day, month, year, absences_data):
    for a in absences_data:
        if a.get("tab_number") == tab and a.get("month") == month and a.get("year") == year:
            if a.get("start_day", 0) <= day <= a.get("end_day", 0):
                return True
    return False


def main():
    input_shifts = config.OUTPUT_SHIFTS
    input_tabel = config.INPUT_TABEL_FILE
    input_absences = config.OUTPUT_ABSENCES_FILE
    output_excel = os.path.join(parent_dir, "output_data", "monthly_report_twice.xlsx")

    if not os.path.exists(input_shifts):
        print(f"Файл не найден: {input_shifts}")
        return

    with open(input_shifts, 'r', encoding='utf-8') as f:
        assigned_shifts = json.load(f)
    with open(input_tabel, 'r', encoding='utf-8') as f:
        tabel_data = json.load(f).get("drivers", [])

    absences_data = []
    if os.path.exists(input_absences):
        with open(input_absences, 'r', encoding='utf-8') as f:
            try:
                absences_data = json.load(f)
            except:
                pass

    planned_schedules = {d["tab_number"]: d.get("days", {}) for d in tabel_data}
    driver_profiles = {d["tab_number"]: d for d in tabel_data}

    driver_history = defaultdict(list)
    for record in assigned_shifts:
        record['start_dt'] = datetime.strptime(record['start_time'], "%Y-%m-%d %H:%M")
        record['end_dt'] = datetime.strptime(record['end_time'], "%Y-%m-%d %H:%M")
        tab = record.get('assigned_driver')
        if tab is not None:
            driver_history[tab].append(record)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Цветной график"

    # --- НАСТРОЙКА ЦВЕТОВ ---
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")  # Синеватый
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Светло-зеленый (Отработано)
    gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")  # Светло-серый (Выходной)
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # Желтый (Резерв)

    bold_font = Font(bold=True)
    center_aligned = Alignment(horizontal="center", vertical="center")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    year = config.TARGET_YEAR
    month = config.TARGET_MONTH
    _, days_in_month = calendar.monthrange(year, month)

    # 1. ФОРМИРОВАНИЕ ШАПКИ
    headers_info = ["Таб. номер", "График", "Режим", "Смен", "Часов"]
    for i, title in enumerate(headers_info, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.alignment = center_aligned
        cell.font = bold_font
        cell.fill = header_fill
        cell.border = border_thin

    for day in range(1, days_in_month + 1):
        cell = ws.cell(row=1, column=5 + day, value=day)
        cell.alignment = center_aligned
        cell.font = bold_font
        cell.fill = header_fill
        cell.border = border_thin
        ws.column_dimensions[openpyxl.utils.get_column_letter(5 + day)].width = 14

    # 2. ЗАПОЛНЕНИЕ ДАННЫХ
    current_row = 2
    daily_reserve_count = {day: 0 for day in range(1, days_in_month + 1)}

    for tab in sorted(driver_history.keys()):
        shifts = sorted(driver_history[tab], key=lambda x: x['start_dt'])
        profile = driver_profiles.get(tab, {})
        total_seconds = sum((s['end_dt'] - s['start_dt']).total_seconds() for s in shifts)

        ws.cell(row=current_row, column=1, value=tab).border = border_thin
        ws.cell(row=current_row, column=2, value=profile.get('schedule_type', '-')).border = border_thin
        ws.cell(row=current_row, column=3, value=profile.get('mode', '-')).border = border_thin
        ws.cell(row=current_row, column=4, value=len(shifts)).border = border_thin
        ws.cell(row=current_row, column=5, value=format_duration(total_seconds)).border = border_thin

        shifts_by_day = {shift['day']: idx for idx, shift in enumerate(shifts)}
        driver_plan = planned_schedules.get(tab, {})

        for day in range(1, days_in_month + 1):
            cell = ws.cell(row=current_row, column=5 + day)
            cell.alignment = center_aligned
            cell.border = border_thin

            if day in shifts_by_day:
                # ЗАКРЫТЫЙ ДЕНЬ (РАБОТА) -> ЗЕЛЕНЫЙ
                idx = shifts_by_day[day]
                curr = shifts[idx]
                work_time = format_duration((curr['end_dt'] - curr['start_dt']).total_seconds())
                cell.fill = green_fill

                if idx + 1 < len(shifts):
                    rest_sec = (shifts[idx + 1]['start_dt'] - curr['end_dt']).total_seconds()
                    rest_time = format_duration(rest_sec)
                    cell.value = f"Работа: {work_time}. Отдых: {rest_time}"
                else:
                    cell.value = f"{work_time} (-)"
            else:
                planned = driver_plan.get(str(day), "В")
                is_absent = is_driver_absent(tab, day, month, year, absences_data)

                if is_absent:
                    cell.value = "Отс"
                elif planned in ["1", "2"]:
                    # РЕЗЕРВ -> ЖЕЛТЫЙ
                    cell.value = "Резерв"
                    cell.fill = yellow_fill
                    daily_reserve_count[day] += 1
                else:
                    # ВЫХОДНОЙ -> СЕРЫЙ
                    cell.value = "В"
                    cell.fill = gray_fill

        current_row += 1

    # 3. СТРОКА ИТОГО
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    total_lbl = ws.cell(row=current_row, column=1, value="ИТОГО РЕЗЕРВ")
    total_lbl.font = bold_font
    total_lbl.alignment = Alignment(horizontal="right", vertical="center")

    for day in range(1, days_in_month + 1):
        res_cell = ws.cell(row=current_row, column=5 + day, value=daily_reserve_count[day])
        res_cell.alignment = center_aligned
        res_cell.font = bold_font
        res_cell.border = border_thin
        if daily_reserve_count[day] > 0:
            res_cell.fill = yellow_fill

    wb.save(output_excel)
    print(f"Отчет с цветовым выделением создан: {output_excel}")


if __name__ == "__main__":
    main()
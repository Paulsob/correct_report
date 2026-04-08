import sys
import os
import json
import calendar
import time
from datetime import date, datetime
import pandas as pd
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

# ВАЖНО: импортируем config из папки blocks_debug
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from blocks_debug import config

# --- НАСТРОЙКИ МОДЕЛИРОВАНИЯ ---
TARGET_SCHEDULES = ["4x2"]  # Для Excel отчета (список)
TARGET_ROUTE = 9  # Укажите номер маршрута (например, 9) или None для всех маршрутов

# Базовые переменные времени
ANCHOR_DATE = 1
YEAR = int(config.TARGET_YEAR)
MONTH = int(config.TARGET_MONTH)
DAYS_IN_MONTH = calendar.monthrange(YEAR, MONTH)[1]
RU_DAYS = {0: "Пн", 1: "Вт", 2: "Ср", 3: "Чт", 4: "Пт", 5: "Сб", 6: "Вс"}

# Определяем пути для сохранения (чтобы файлы содержали номер маршрута)
OUT_DIR = os.path.dirname(config.FINAL_SCHEDULE)  # Берем папку из конфига
ROUTE_SUFFIX = f"_route_{TARGET_ROUTE}" if TARGET_ROUTE is not None else ""
CUSTOM_JSON_PATH = os.path.join(OUT_DIR, f"DEBUG_{config.TARGET_SCHEDULE}{ROUTE_SUFFIX}_schedule.json")
CUSTOM_EXCEL_PATH = os.path.join(OUT_DIR, f"DEBUG_{config.TARGET_SCHEDULE}{ROUTE_SUFFIX}_report.xlsx")


# ==========================================
# ЧАСТЬ 1: ГЕНЕРАЦИЯ РАСПИСАНИЯ
# ==========================================

def load_matrix_data(schedule_type):
    with open(config.MATRICES_FILE, "r", encoding="utf-8") as f:
        all_matrices = json.load(f)
    if schedule_type not in all_matrices:
        raise ValueError(f"Ошибка: График '{schedule_type}' не найден в {config.MATRICES_FILE}!")
    matrix_data = all_matrices[schedule_type]
    active_matrix = {int(k): v for k, v in matrix_data["matrix"].items()}
    return active_matrix, matrix_data["cycle_length"]


def get_required_role(cycle_day, target_slot, active_matrix):
    day_index = cycle_day - 1
    for role, days_list in active_matrix.items():
        if days_list[day_index] == target_slot:
            return role
    return None


def build_drivers_lookup(drivers_data, schedule_type):
    lookup = {}
    for driver in drivers_data.get("drivers", []):
        if driver.get("schedule_type") != schedule_type:
            continue
        b_id = driver.get("block_id")
        m_role = driver.get("matrix_role")
        if b_id not in lookup:
            lookup[b_id] = {}
        driver["schedule_dict"] = driver.get("days", {})
        lookup[b_id][m_role] = driver
    return lookup


def get_day_type(day):
    # ДЕБАГ-РЕЖИМ: Все дни считаются рабочими.
    return "рабочий"


def generate_schedule():
    start_time = time.time()

    print(f"\n{'=' * 60}")
    print(f"🔧 СТАРТ МОДЕЛИРОВАНИЯ (DEBUG РЕЖИМ) 🔧")
    print(f"График: [{config.TARGET_SCHEDULE}] | Период: {MONTH:02d}.{YEAR}")
    route_display = f"ТОЛЬКО МАРШРУТ №{TARGET_ROUTE}" if TARGET_ROUTE else "ВСЕ МАРШРУТЫ"
    print(f"Режим: ВСЕ ДНИ РАБОЧИЕ | Фильтр: {route_display}")
    print(f"{'=' * 60}")

    print("\n[1] ИСТОЧНИКИ ДАННЫХ:")
    print(f"  -> Исходное расписание: {config.SCHEDULE_PREPARED}")
    print(f"  -> Матрицы графиков: {config.MATRICES_FILE}")
    print(f"  -> База водителей: {config.DRIVERS_PREPARED}")

    active_matrix, cycle_length = load_matrix_data(config.TARGET_SCHEDULE)

    with open(config.SCHEDULE_PREPARED, "r", encoding="utf-8") as f:
        trams_data = json.load(f)

    with open(config.DRIVERS_PREPARED, "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    drivers_lookup = build_drivers_lookup(drivers_data, config.TARGET_SCHEDULE)

    monthly_result = []

    print("\n[2] ГЕНЕРАЦИЯ РАСПИСАНИЯ ПО ДНЯМ...")
    for current_day in range(1, DAYS_IN_MONTH + 1):
        cycle_day = ((current_day - ANCHOR_DATE) % cycle_length) + 1
        current_day_type = get_day_type(current_day)

        daily_assignments = []
        daily_unassigned = []

        for route in trams_data:
            # 1. ФИЛЬТР ПО ТИПУ ДНЯ
            if route.get("день") and route.get("день") != current_day_type:
                continue

            # 2. ФИЛЬТР ПО МАРШРУТУ
            route_num = route.get("маршрут")
            if TARGET_ROUTE is not None and str(route_num) != str(TARGET_ROUTE):
                continue

            for tram in route.get("трамваи", []):
                block_id = tram.get("block_ids", {}).get(config.TARGET_SCHEDULE)
                if not block_id:
                    continue

                tram_num = tram.get("номер")
                shifts = []
                if tram.get("смена_1"): shifts.append(("Утро", tram["смена_1"]))
                if tram.get("смена_2"): shifts.append(("Вечер", tram["смена_2"]))

                for shift_name, shift_data in shifts:
                    target_slot = shift_data.get("matrix_slots", {}).get(config.TARGET_SCHEDULE)
                    if not target_slot:
                        continue

                    required_role = get_required_role(cycle_day, target_slot, active_matrix)
                    if not required_role:
                        continue

                    driver = drivers_lookup.get(block_id, {}).get(required_role)

                    if not driver:
                        daily_unassigned.append({
                            "маршрут": route_num,
                            "трамвай": tram_num,
                            "смена": shift_name,
                            "нужна_роль": required_role,
                            "причина": f"В блоке {block_id} нет водителя с ролью {required_role}"
                        })
                        continue

                    expected_status = "1" if target_slot % 2 != 0 else "2"
                    actual_status = driver["schedule_dict"].get(str(current_day), "Нет данных")

                    if actual_status == expected_status:
                        daily_assignments.append({
                            "маршрут": route_num,
                            "трамвай": tram_num,
                            "смена": shift_name,
                            "время": f"{shift_data.get('отправление')} - {shift_data.get('прибытие')}",
                            "таб_номер_водителя": driver.get("tab_number"),
                            "роль_по_матрице": required_role
                        })
                    else:
                        daily_unassigned.append({
                            "маршрут": route_num,
                            "трамвай": tram_num,
                            "смена": shift_name,
                            "нужна_роль": required_role,
                            "таб_номер_водителя": driver.get("tab_number"),
                            "причина": f"Ожидался статус '{expected_status}', а в графике '{actual_status}'"
                        })

        monthly_result.append({
            "дата": f"{YEAR}-{MONTH:02d}-{current_day:02d}",
            "тип_дня": current_day_type,
            "день_цикла": cycle_day,
            "успешные_назначения": daily_assignments,
            "открытые_смены_без_водителя": daily_unassigned
        })

    with open(CUSTOM_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(monthly_result, f, ensure_ascii=False, indent=2)

    execution_time = time.time() - start_time
    print(f"[OK] Расписание сгенерировано за {execution_time:.2f} сек.")
    print(f"Сохранено в: {CUSTOM_JSON_PATH}")


# ==========================================
# ЧАСТЬ 2: ГЕНЕРАЦИЯ EXCEL ОТЧЕТА
# ==========================================

def calculate_hours(time_range_str):
    try:
        start_str, end_str = time_range_str.split(" - ")
        time_format = "%H:%M"
        start_time = datetime.strptime(start_str, time_format)
        end_time = datetime.strptime(end_str, time_format)

        delta_hours = (end_time - start_time).total_seconds() / 3600.0
        if delta_hours < 0:
            delta_hours += 24.0
        return round(delta_hours, 2)
    except Exception as e:
        return 0.0


def generate_excel():
    print(f"\n{'=' * 60}")
    print("📊 ФОРМИРОВАНИЕ EXCEL ОТЧЕТА 📊")
    print(f"{'=' * 60}")
    print(f"  -> Чтение готовых смен из: {CUSTOM_JSON_PATH}")

    if not os.path.exists(CUSTOM_JSON_PATH):
        print(f"Ошибка: Файл расписания {CUSTOM_JSON_PATH} не найден!")
        return

    with open(CUSTOM_JSON_PATH, "r", encoding="utf-8") as f:
        schedule_data = json.load(f)

    with open(config.DRIVERS_PREPARED, "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    drivers_stats = {}
    for driver in drivers_data.get("drivers", []):
        sched_type = driver.get("schedule_type")
        if sched_type not in TARGET_SCHEDULES:
            continue

        tab = driver.get("tab_number")
        drivers_stats[tab] = {
            "График": sched_type,
            "Табельный номер": tab,
            "Смен в месяц": 0,
            "Часов в месяц": 0.0
        }

        for day_str, val in driver.get("days", {}).items():
            if val == "В":
                drivers_stats[tab][day_str] = "Выходной"
            else:
                drivers_stats[tab][day_str] = "РЕЗЕРВ"

    max_day_in_month = 0

    for day_info in schedule_data:
        date_str = day_info.get("дата", "")
        if not date_str:
            continue

        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        day_number = date_obj.day
        if day_number > max_day_in_month:
            max_day_in_month = day_number

        for assignment in day_info.get("успешные_назначения", []):
            tab_number = assignment.get("таб_номер_водителя")
            shift_type = assignment.get("смена")
            time_str = assignment.get("время")
            route_num = assignment.get("маршрут")
            tram_num = assignment.get("трамвай")

            if tab_number not in drivers_stats:
                continue

            hours = calculate_hours(time_str)
            drivers_stats[tab_number]["Смен в месяц"] += 1
            drivers_stats[tab_number]["Часов в месяц"] += hours

            cell_text = (f"Маршрут: {route_num}\n"
                         f"Режим: {shift_type}\n"
                         f"Смена №: {tram_num}\n"
                         f"{hours} часов")

            current_val = drivers_stats[tab_number][str(day_number)]
            if current_val in ["РЕЗЕРВ", "Выходной"]:
                drivers_stats[tab_number][str(day_number)] = cell_text
            else:
                drivers_stats[tab_number][str(day_number)] += f"\n\n{cell_text}"

    # ФИЛЬТРАЦИЯ: Оставляем только тех, кто РЕАЛЬНО работал на ЭТОМ маршруте
    drivers_stats = {tab: data for tab, data in drivers_stats.items() if data["Смен в месяц"] > 0}

    data_for_df = list(drivers_stats.values())
    if not data_for_df:
        print(f"Нет ни одного водителя с назначенными сменами на маршруте {TARGET_ROUTE}. Отчет пуст.")
        return

    df = pd.DataFrame(data_for_df)
    df = df.sort_values(by=["График", "Табельный номер"])
    df["Часов в месяц"] = df["Часов в месяц"].round(2)

    # --- ДОБАВЛЯЕМ СТРОКУ "ИТОГО РЕЗЕРВА" ВНИЗ ТАБЛИЦЫ ---
    total_row = {
        "График": "ИТОГО:",
        "Табельный номер": "Человек в резерве ->",
        "Смен в месяц": "",
        "Часов в месяц": ""
    }

    # Считаем количество слова "РЕЗЕРВ" в каждой колонке-дне
    for d in range(1, max_day_in_month + 1):
        day_str = str(d)
        count_reserve = (df[day_str] == "РЕЗЕРВ").sum()
        total_row[day_str] = count_reserve

    # Добавляем строку с итогами в конец DataFrame
    df_total = pd.DataFrame([total_row])
    df = pd.concat([df, df_total], ignore_index=True)
    # ----------------------------------------------------

    new_columns = []
    columns_to_keep = ["График", "Табельный номер", "Смен в месяц", "Часов в месяц"] + [str(d) for d in
                                                                                        range(1, max_day_in_month + 1)]
    df = df[columns_to_keep]

    for col in df.columns:
        if col in ["График", "Табельный номер", "Смен в месяц", "Часов в месяц"]:
            new_columns.append(col)
        else:
            day = int(col)
            current_date = datetime(YEAR, MONTH, day)
            day_of_week_str = RU_DAYS[current_date.weekday()]
            new_columns.append(f"{day}\n{day_of_week_str}")

    df.columns = new_columns

    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    gray_fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    total_fill = PatternFill(start_color="FFD9EAD3", end_color="FFD9EAD3",
                             fill_type="solid")  # Нежно-зеленый для итогов
    center_aligned_text = Alignment(horizontal="center", vertical="center", wrap_text=True)

    with pd.ExcelWriter(CUSTOM_EXCEL_PATH, engine="openpyxl") as writer:
        sheet_name = f"Отчет М{TARGET_ROUTE}"[:31] if TARGET_ROUTE else "Отчет Все"
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        worksheet.freeze_panes = "E2"
        worksheet.row_dimensions[1].height = 30

        # Высота строк: для обычных водителей - широкая, для строки ИТОГО (последняя) - узкая
        for r in range(2, worksheet.max_row):
            worksheet.row_dimensions[r].height = 65
        worksheet.row_dimensions[worksheet.max_row].height = 25  # Строка ИТОГО

        for idx, col in enumerate(df.columns):
            col_letter = get_column_letter(idx + 1)
            worksheet.column_dimensions[col_letter].width = 12 if idx < 4 else 16

        # Стилизация основных ячеек
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row - 1, min_col=5,
                                       max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = center_aligned_text
                if cell.value == "РЕЗЕРВ":
                    cell.fill = yellow_fill
                elif cell.value == "Выходной":
                    cell.fill = gray_fill
                    cell.font = Font(color="000000")

        # Центрирование текста в первых колонках
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row - 1, min_col=1, max_col=4):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Стилизация шапки таблицы
        for cell in worksheet[1]:
            cell.alignment = center_aligned_text
            cell.font = Font(bold=True)

        # --- Стилизация добавленной строки ИТОГО (последняя строка) ---
        for cell in worksheet[worksheet.max_row]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(bold=True)
            cell.fill = total_fill

    print(f"[OK] Отчет успешно сформирован!")
    print(f"Сохранено в: {CUSTOM_EXCEL_PATH}\n")


if __name__ == "__main__":
    generate_schedule()
    generate_excel()
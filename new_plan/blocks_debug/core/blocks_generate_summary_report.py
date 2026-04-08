import sys
import os

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from blocks_debug import config

import json
import pandas as pd
from datetime import datetime
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.utils import get_column_letter

RU_DAYS = {0: "Пн", 1: "Вт", 2: "Ср", 3: "Чт", 4: "Пт", 5: "Сб", 6: "Вс"}

# --- НАСТРОЙКИ ОТЧЕТА ---
TARGET_SCHEDULES = ["4x2"]


# TARGET_SCHEDULES = ["4x2"] # <-- для старого моно-режима

def calculate_hours(time_range_str):
    try:
        start_str, end_str = time_range_str.split(" - ")
        time_format = "%H:%M"
        start_time = datetime.strptime(start_str, time_format)
        end_time = datetime.strptime(end_str, time_format)

        delta_hours = (end_time - start_time).total_seconds() / 3600.0
        if delta_hours < 0:
            delta_hours += 24.0

        return round(delta_hours, 2)
    except Exception as e:
        print(f"Ошибка парсинга времени '{time_range_str}': {e}")
        return 0.0


def generate_excel():
    prefix = "_".join(TARGET_SCHEDULES)
    print(f"Загрузка данных для формирования отчета: [{prefix}]...")

    # Формируем динамические пути (как в планировщике)
    YEAR = int(config.TARGET_YEAR)
    MONTH = int(config.TARGET_MONTH)
    blocks_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_dir = os.path.join(blocks_dir, "output_data", str(YEAR), f"{MONTH:02d}")

    json_filename = f"{prefix}_final_schedule.json"
    excel_filename = f"{prefix}_drivers_report.xlsx"

    schedule_path = os.path.join(output_dir, json_filename)
    report_path = os.path.join(output_dir, excel_filename)

    if not os.path.exists(schedule_path) or not os.path.exists(config.DRIVERS_PREPARED):
        print(f"Ошибка: Файл расписания {schedule_path} не найден! Сначала запустите планировщик.")
        return

    with open(schedule_path, "r", encoding="utf-8") as f:
        schedule_data = json.load(f)

    with open(config.DRIVERS_PREPARED, "r", encoding="utf-8") as f:
        drivers_data = json.load(f)

    # 1. Формируем "Подложку" для всех графиков из списка
    drivers_stats = {}
    for driver in drivers_data.get("drivers", []):
        sched_type = driver.get("schedule_type")

        if sched_type not in TARGET_SCHEDULES:
            continue

        tab = driver.get("tab_number")
        drivers_stats[tab] = {
            "График": sched_type,  # <-- Добавили колонку для визуального разделения
            "Табельный номер": tab,
            "Смен в месяц": 0,
            "Часов в месяц": 0.0
        }

        # Читаем сжатый словарь дней ("1": "1", "2": "В")
        for day_str, val in driver.get("days", {}).items():
            if val == "В":
                drivers_stats[tab][day_str] = "Выходной"
            else:
                drivers_stats[tab][day_str] = "РЕЗЕРВ"

    max_day_in_month = 0

    # 2. Накладываем фактические назначения
    print("Обработка назначений на смены...")
    for day_info in schedule_data:
        date_str = day_info.get("дата", "")
        if not date_str:
            continue

        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        day_number = date_obj.day

        if day_number > max_day_in_month:
            max_day_in_month = day_number

        for assignment in day_info.get("успешные_назначения", []):
            tab_number = assignment.get("таб_номер_водителя")
            shift_type = assignment.get("смена")
            time_str = assignment.get("время")
            route_num = assignment.get("маршрут")
            tram_num = assignment.get("трамвай")

            if tab_number not in drivers_stats:
                continue

            hours = calculate_hours(time_str)

            drivers_stats[tab_number]["Смен в месяц"] += 1
            drivers_stats[tab_number]["Часов в месяц"] += hours

            cell_text = (f"Маршрут: {route_num}\n"
                         f"Режим: {shift_type}\n"
                         f"Смена №: {tram_num}\n"
                         f"{hours} часов")

            current_val = drivers_stats[tab_number][str(day_number)]
            if current_val in ["РЕЗЕРВ", "Выходной"]:
                drivers_stats[tab_number][str(day_number)] = cell_text
            else:
                drivers_stats[tab_number][str(day_number)] += f"\n\n{cell_text}"

    # 3. ФИЛЬТРАЦИЯ ВОДИТЕЛЕЙ (Оставляем только тех, кто работал)
    drivers_stats = {tab: data for tab, data in drivers_stats.items() if data["Смен в месяц"] > 0}

    data_for_df = list(drivers_stats.values())
    if not data_for_df:
        print(f"Нет ни одного водителя с назначенными сменами. Отчет пуст.")
        return

    df = pd.DataFrame(data_for_df)

    # Сортируем сначала по Графику, потом по Табельному номеру
    df = df.sort_values(by=["График", "Табельный номер"])
    df["Часов в месяц"] = df["Часов в месяц"].round(2)

    # 4. Создаем умную шапку
    new_columns = []
    columns_to_keep = ["График", "Табельный номер", "Смен в месяц", "Часов в месяц"] + [str(d) for d in
                                                                                        range(1, max_day_in_month + 1)]
    df = df[columns_to_keep]

    for col in df.columns:
        if col in ["График", "Табельный номер", "Смен в месяц", "Часов в месяц"]:
            new_columns.append(col)
        else:
            day = int(col)
            current_date = datetime(YEAR, MONTH, day)
            day_of_week_str = RU_DAYS[current_date.weekday()]
            new_columns.append(f"{day}\n{day_of_week_str}")

    df.columns = new_columns

    # 5. Сохранение и визуальное оформление
    print("Сохранение и форматирование в Excel...")

    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    gray_fill = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
    center_aligned_text = Alignment(horizontal="center", vertical="center", wrap_text=True)

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        sheet_name = f"Отчет {prefix}"[:31]  # Имя листа в Excel ограничено 31 символом
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        worksheet.freeze_panes = "E2"  # Смещаем закрепление из-за новой колонки "График"
        worksheet.row_dimensions[1].height = 30

        for r in range(2, worksheet.max_row + 1):
            worksheet.row_dimensions[r].height = 65

        for idx, col in enumerate(df.columns):
            col_letter = get_column_letter(idx + 1)
            # Колонки "График" и "Таб. номер" делаем уже, дни - шире
            worksheet.column_dimensions[col_letter].width = 12 if idx < 4 else 16

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=5, max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = center_aligned_text
                if cell.value == "РЕЗЕРВ":
                    cell.fill = yellow_fill
                elif cell.value == "Выходной":
                    cell.fill = gray_fill
                    cell.font = Font(color="000000")

        # Центрируем всё в колонках слева (График, Табельный и т.д.)
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for cell in worksheet[1]:
            cell.alignment = center_aligned_text
            cell.font = Font(bold=True)

    print(f"Отчет успешно сформирован и сохранен в:\n{report_path}")


if __name__ == "__main__":
    generate_excel()